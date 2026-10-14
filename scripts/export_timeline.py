# -*- coding: utf-8 -*-
"""
导出采样时间表
================================================================================
从 data/multi_boards.db 的 heat_samples 提取全部采样时刻，产出两张表：

  Sheet1「采样时刻」    —— 每一次采样的时刻、偏移、间隔、样本数
  Sheet2「模板对照」    —— 课程要求的 1/2/4/6/8/10/12/18/24/36/48/96 小时
                          相对首个采样点的实际覆盖情况

用法：
  python -u scripts/export_timeline.py
输出：
  dataset_boards/六榜数据集/采样时间表.xlsx
"""
import os
import sqlite3
import sys
from datetime import datetime, timedelta

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DB = os.path.join(ROOT, 'data', 'multi_boards.db')
OUT = os.path.join(ROOT, 'dataset_boards', '六榜数据集')

WEEK = ['周一', '周二', '周三', '周四', '周五', '周六', '周日']
TPL_HOURS = [1, 2, 4, 6, 8, 10, 12, 18, 24, 36, 48, 96]

HEAD_FILL = PatternFill('solid', fgColor='2F5597')
HEAD_FONT = Font(color='FFFFFF', bold=True, size=10)
OK_FILL = PatternFill('solid', fgColor='E2EFDA')
BAD_FILL = PatternFill('solid', fgColor='FCE4E4')
THIN = Side(style='thin', color='BFBFBF')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def style(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for cell in ws[1]:
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = BORDER
    ws.freeze_panes = 'A2'
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.border = BORDER
            cell.alignment = Alignment(horizontal='center', vertical='center')


def main():
    if not os.path.exists(DB):
        print(f'!! 找不到 {DB}')
        return 1
    os.makedirs(OUT, exist_ok=True)

    c = sqlite3.connect(DB)
    rows = c.execute('SELECT sample_ts, COUNT(*), COUNT(DISTINCT topic_id) '
                     'FROM heat_samples GROUP BY sample_ts ORDER BY sample_ts').fetchall()
    if not rows:
        print('!! heat_samples 为空')
        return 1

    ticks = [r[0] for r in rows]
    t0 = datetime.strptime(ticks[0], '%Y-%m-%d %H:%M:%S')

    wb = Workbook()

    # ---------- Sheet1：采样时刻 ----------
    ws = wb.active
    ws.title = '采样时刻'
    ws.append(['序号', '采样时刻', '星期', '距首个采样点(h)', '距上次间隔(min)',
               '热度样本数', '话题数'])
    prev = None
    for i, (ts, n, ntopic) in enumerate(rows, 1):
        t = datetime.strptime(ts, '%Y-%m-%d %H:%M:%S')
        off = round((t - t0).total_seconds() / 3600, 2)
        gap = '' if prev is None else int(round((t - prev).total_seconds() / 60))
        ws.append([i, ts, WEEK[t.weekday()], off, gap, n, ntopic])
        prev = t
    style(ws, [6, 21, 8, 17, 18, 12, 10])

    # ---------- Sheet2：模板对照 ----------
    ws2 = wb.create_sheet('模板对照')
    ws2.append(['要求偏移', '目标时刻', '最近采样点', '实际偏移(h)', '偏差(min)', '判定'])
    dts = [datetime.strptime(x, '%Y-%m-%d %H:%M:%S') for x in ticks]
    for h in TPL_HOURS:
        tgt = t0 + timedelta(hours=h)
        near = min(dts, key=lambda t: abs((t - tgt).total_seconds()))
        err = int(round((near - tgt).total_seconds() / 60))
        act = round((near - t0).total_seconds() / 3600, 2)
        ok = '✓ 覆盖' if abs(err) <= 30 else ('△ 勉强' if abs(err) <= 60 else '✗ 缺口')
        ws2.append([f'{h}h', tgt.strftime('%Y-%m-%d %H:%M:%S'),
                    near.strftime('%Y-%m-%d %H:%M:%S'), act, err, ok])
    style(ws2, [10, 21, 21, 14, 12, 10])
    for r in range(2, ws2.max_row + 1):
        fill = OK_FILL if str(ws2.cell(r, 6).value).startswith('✓') else BAD_FILL
        for col in range(1, 7):
            ws2.cell(r, col).fill = fill

    path = os.path.join(OUT, '采样时间表.xlsx')
    wb.save(path)
    print(f'✓ {path}')
    print(f'  采样时刻 {len(rows)} 个 · 跨度 {(dts[-1] - t0).total_seconds() / 3600:.1f} 小时')
    n_ok = sum(1 for h in TPL_HOURS
               if abs((min(dts, key=lambda t: abs((t - (t0 + timedelta(hours=h))).total_seconds()))
                       - (t0 + timedelta(hours=h))).total_seconds() / 60) <= 30)
    print(f'  模板 12 点覆盖：{n_ok}/12')
    return 0


if __name__ == '__main__':
    sys.exit(main())
