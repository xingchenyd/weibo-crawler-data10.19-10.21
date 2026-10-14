# -*- coding: utf-8 -*-
"""
回填「图片」列到 事件表.xlsx
================================================================================
与 fetch_boards_pics.py 的区别：本脚本**只读磁盘上已存在的图片文件**，
不发起任何网络请求，因此可以随时重复运行、秒级完成。

用途：fetch_boards_pics.py 下载完图片后若因「事件表.xlsx 被 WPS/Excel 打开」
      导致 wb.save() 报 PermissionError，可用本脚本补写那一列。

用法：
  python -u scripts/fill_pic_column.py
输出：
  dataset_boards/六榜数据集/事件表.xlsx（新增/更新「图片」列）
"""
import os
import sys
from collections import defaultdict

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUTDIR = os.path.join(ROOT, 'dataset_boards', '六榜数据集')
IMGDIR = os.path.join(OUTDIR, '事件图')
XLSX = os.path.join(OUTDIR, '事件表.xlsx')

HEAD_FILL = PatternFill('solid', fgColor='2F5597')
HEAD_FONT = Font(color='FFFFFF', bold=True, size=10)
THIN = Side(style='thin', color='BFBFBF')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def main():
    if not os.path.exists(XLSX):
        print(f'!! 找不到 {XLSX}')
        return 1
    if not os.path.isdir(IMGDIR):
        print(f'!! 找不到图片目录 {IMGDIR}')
        return 1

    # ---------- 扫描磁盘：序号 -> 图片张数 ----------
    have = defaultdict(int)
    for fn in os.listdir(IMGDIR):
        if not fn.lower().endswith('.jpg'):
            continue
        no = fn.split('-', 1)[0]
        if no.isdigit():
            have[int(no)] += 1
    print(f'磁盘上共 {sum(have.values())} 张图，覆盖 {len(have)} 个事件')

    # ---------- 写回表格 ----------
    try:
        wb = openpyxl.load_workbook(XLSX)
    except PermissionError:
        print('!! 事件表.xlsx 正被 WPS/Excel 打开，无法写入。')
        print('   请先关闭该文件，再重新运行本脚本。')
        return 2

    ws = wb.active
    hdr = [c.value for c in ws[1]]
    if '图片' in hdr:
        ci = hdr.index('图片')
        print('（已存在「图片」列，将覆盖更新）')
    else:
        ci = len(hdr)
        cell = ws.cell(row=1, column=ci + 1, value='图片')
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = BORDER
        print('（原表无「图片」列，已追加到最后一列）')

    ci_no = hdr.index('序号') if '序号' in hdr else 0
    n_fill, n_empty = 0, 0
    for r in range(2, ws.max_row + 1):
        no = ws.cell(row=r, column=ci_no + 1).value
        if isinstance(no, int) and have.get(no):
            ws.cell(row=r, column=ci + 1, value=f'P{no}')
            n_fill += 1
        else:
            n_empty += 1

    try:
        wb.save(XLSX)
    except PermissionError:
        print('!! 保存失败：事件表.xlsx 仍被占用，请关闭后重试。')
        return 2

    print(f'✓ 已写入 {XLSX}')
    print(f'  填了图片引用 {n_fill} 行 · 无图 {n_empty} 行')
    print('  「图片」列填 P{序号}，对应文件 事件图\\{序号}-{k}.jpg')
    return 0


if __name__ == '__main__':
    sys.exit(main())
